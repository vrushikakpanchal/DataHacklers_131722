import os
import logging

logging.getLogger('weasyprint').setLevel(logging.ERROR)
logging.getLogger('fontTools').setLevel(logging.ERROR)
logging.getLogger('glib').setLevel(logging.ERROR)
import json
import re
import ollama
from html import escape
from textwrap import wrap
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from typing import Dict, Any
from modules.design_engine import as_items, infer_design, svg_icon

# Ensure output directory exists
OUTPUT_DIR = os.path.join("data", "outputs")
os.makedirs(OUTPUT_DIR, exist_ok=True)


# ---------------------------------------------------------
# Helper: Extract Canonical Facts Dictionary from Raw Text
# ---------------------------------------------------------
def build_canonical_facts(
    text: str,
    locked_params: Dict[str, Any] = None,
    model_name: str = "qwen2.5:3b",
    language: str = "English",
    detail_level: str = "Executive Brief",
    objective: str = "Inform",
    content_style: str = "Corporate",
) -> Dict[str, Any]:
    """
    Parses unstructured advisory text into a structured facts dictionary.
    Integrates locked deterministic parameters (CVEs, IPs, Severities) if provided.
    """
    prompt = f"""
    Analyze the following advisory text and extract structured facts.
    Write human-readable values in {language}, at a {detail_level} level, using a {content_style} style for the objective: {objective}.
    Preserve all source facts and do not translate or alter technical identifiers.
    
    Source Text:
    {text}
    
    Return a strictly valid JSON object with the following schema:
    {{
        "title": "Short descriptive advisory title",
        "severity": "CRITICAL / HIGH / MEDIUM / LOW",
        "summary": "2-3 sentence executive summary",
        "cve_ids": ["CVE-YYYY-XXXX"],
        "affected_systems": ["Every platform, product, service, website, app, software, hardware model, or named entity that is explicitly mentioned in the source as affected, targeted, misused, or involved — e.g. Shaadi.com, Tinder, WhatsApp, UPI, Apache HTTP Server. Include ALL such names found in the text. Leave empty ONLY if no such names appear."],
        "locked_ips": ["IP addresses mentioned"],
        "recommended_actions": ["Specific actionable mitigation step 1", "Step 2"]
    }}
    
    IMPORTANT for affected_systems: scan the entire source text for any named platforms, apps, websites, services, software products, hardware models, or organizations described as affected or involved. Do not restrict to technical software — include consumer apps, social networks, payment services, websites, and any explicitly named entity.
    
    Do NOT include markdown wrapping or extra text outside JSON.
    """
    
    try:
        response = ollama.chat(
            model=model_name,
            messages=[{'role': 'user', 'content': prompt}],
            format='json'
        )
        facts = json.loads(response['message']['content'])
    except Exception:
        facts = {
            "title": "Security Advisory Briefing",
            "severity": "HIGH",
            "summary": text[:250] + "...",
            "cve_ids": [],
            "affected_systems": [],
            "locked_ips": [],
            "recommended_actions": ["Review log files and apply latest system patches."]
        }

    # Merge locked deterministic parameters if available from extractor
    if locked_params:
        if locked_params.get("locked_cves"):
            facts["cve_ids"] = list(set(facts.get("cve_ids", []) + locked_params["locked_cves"]))
        if locked_params.get("locked_ips"):
            facts["locked_ips"] = list(set(facts.get("locked_ips", []) + locked_params["locked_ips"]))
        if locked_params.get("locked_severities") and not facts.get("severity"):
            facts["severity"] = locked_params["locked_severities"][0]
        # If the LLM returned no affected systems, fall back to deterministically
        # extracted entity names from the source text so the list is never empty
        # when named platforms/products are present.
        if not facts.get("affected_systems") and locked_params.get("locked_entities"):
            facts["affected_systems"] = list(locked_params["locked_entities"])

    return facts


# ---------------------------------------------------------
# 1. Text Transformation Generators for Streamlit App
# ---------------------------------------------------------
def generate_summary(
    text: str,
    model_name: str = "qwen2.5:3b",
    language: str = "English",
    detail_level: str = "Executive Brief",
    objective: str = "Inform",
    content_style: str = "Corporate",
) -> str:
    """Generate an executive summary using operator-selected output settings."""
    system_prompt = (
        "You are a technical content summarizer. Rules:\n"
        "1. Every statement must be directly supported by the provided source text; never add outside knowledge or assumptions.\n"
        "2. Never emit placeholders or mock content: no bracketed slots such as [insert X] or <TBD>, no 'Lorem ipsum', "
        "no example or sample figures such as 'XX%', and no invented numbers, percentages, dates, CVE IDs, IP addresses, "
        "product names, vendor claims, or statistics. If the source does not state something, omit it instead of guessing.\n"
        "3. Reproduce all technical identifiers (CVE IDs, IP addresses, product names, version numbers) exactly as written in the source.\n"
        "4. Output only the requested summary content; do not mention these rules, the source text, or the process."
    )
    prompt = (
        f"Write in {language}. Detail level: {detail_level}. Objective: {objective}. "
        f"Content style: {content_style}. Provide an executive summary of key findings and impact "
        f"using only the source text below.\n\nSOURCE TEXT:\n{text}"
    )
    response = ollama.chat(
        model=model_name,
        messages=[
            {'role': 'system', 'content': system_prompt},
            {'role': 'user', 'content': prompt},
        ],
    )
    return response['message']['content']


def generate_faqs(
    text: str,
    model_name: str = "qwen2.5:3b",
    language: str = "English",
    detail_level: str = "Executive Brief",
    objective: str = "Inform",
    content_style: str = "Corporate",
) -> str:
    """Generate key questions and answers using operator-selected settings."""
    prompt = (
        f"Write in {language}. Detail level: {detail_level}. Objective: {objective}. "
        f"Content style: {content_style}. Preserve technical facts and identifiers. "
        f"Generate 4 key Questions & Answers explaining the core issues and mitigations "
        f"in this text:\n\n{text}"
    )
    response = ollama.chat(model=model_name, messages=[{'role': 'user', 'content': prompt}])
    return response['message']['content']


def generate_report(
    text: str,
    model_name: str = "qwen2.5:3b",
    language: str = "English",
    detail_level: str = "Detailed",
    objective: str = "Mitigate",
    content_style: str = "Technical",
) -> str:
    """Generate a technical advisory report using operator-selected settings."""
    prompt = (
        f"Write in {language}. Detail level: {detail_level}. Objective: {objective}. "
        f"Content style: {content_style}. Preserve technical facts and identifiers. "
        "Construct a security report with sections for Overview, Threat Vectors, "
        "Impact Analysis, and Immediate Recommendations:\n\n" + text
    )
    response = ollama.chat(model=model_name, messages=[{'role': 'user', 'content': prompt}])
    return response['message']['content']


# ---------------------------------------------------------
# 2. Executive Brief & PDF Advisory Generator
# ---------------------------------------------------------
_SEVERITY_COLORS = {"CRITICAL": "#B42332", "HIGH": "#D14B3F", "MEDIUM": "#D28A25", "LOW": "#34866D"}


def _fact_items(value: Any) -> list[str]:
    """Normalize an optional fact field into a list of non-empty strings."""
    if value is None:
        return []
    if isinstance(value, (list, tuple, set)):
        return [str(item).strip() for item in value if str(item).strip()]
    text = str(value).strip()
    return [text] if text else []


def _fact_headline(facts: Dict[str, Any], limit: int = 90) -> str:
    """Derive a document headline strictly from supplied facts."""
    for key in ("title", "topic", "category", "content_type"):
        items = _fact_items(facts.get(key))
        if items:
            return items[0][:limit].rstrip()
    summary = " ".join(_fact_items(facts.get("summary")))
    if summary:
        if len(summary) > limit:
            return summary[:limit].rsplit(" ", 1)[0].rstrip(" ,;:-") + "…"
        return summary
    return "Advisory"


def generate_pdf_advisory(canonical_facts: Dict[str, Any]) -> str:
    """Render the advisory as HTML/PDF using only the supplied canonical facts."""
    facts = canonical_facts or {}
    title = escape(_fact_headline(facts))
    summary = " ".join(_fact_items(facts.get("summary")))
    cves = _fact_items(facts.get("cve_ids") or facts.get("cve_id"))
    systems = _fact_items(facts.get("affected_systems") or facts.get("affected_products"))
    ips = _fact_items(facts.get("locked_ips"))
    actions = _fact_items(
        facts.get("recommended_actions") or facts.get("recommendations") or facts.get("mitigation_steps") or facts.get("key_points")
    )
    severities = _fact_items(facts.get("severity") or facts.get("risk_level"))
    severity = severities[0].upper() if severities else ""

    severity_html = ""
    if severity:
        badge_color = _SEVERITY_COLORS.get(severity, "#4A5568")
        severity_html = f'<p><strong>Severity:</strong> <span class="badge" style="background-color: {badge_color};">{escape(severity)}</span></p>'

    summary_section = ""
    if summary:
        summary_section = f"""
        <div class="section">
            <div class="section-title">Executive Summary</div>
            <p>{escape(summary)}</p>
        </div>"""

    indicator_rows = []
    if cves:
        indicator_rows.append(f'<p><strong>Associated CVEs:</strong> {escape(", ".join(cves))}</p>')
    if systems:
        indicator_rows.append(f'<p><strong>Affected Systems:</strong> {escape(", ".join(systems))}</p>')
    if ips:
        indicator_rows.append(f'<p><strong>Locked IP Addresses:</strong></p><div class="code-box">{escape(", ".join(ips))}</div>')
    indicators_section = ""
    if indicator_rows:
        indicators_section = f"""
        <div class="section">
            <div class="section-title">Identified Vulnerabilities & Technical Indicators</div>
            {"".join(indicator_rows)}
        </div>"""

    mitigations_section = ""
    if actions:
        mitigations_section = f"""
        <div class="section">
            <div class="section-title">Recommended Mitigations</div>
            <ul>{''.join(f'<li>{escape(str(action))}</li>' for action in actions)}</ul>
        </div>"""

    html_content = f"""
    <!DOCTYPE html>
    <html>
    <head>
        <style>
            body {{ font-family: Arial, sans-serif; margin: 40px; color: #222; }}
            .header {{ border-bottom: 3px solid #1a365d; padding-bottom: 10px; margin-bottom: 20px; }}
            .title {{ color: #1a365d; font-size: 24px; font-weight: bold; }}
            .badge {{ display: inline-block; padding: 4px 12px; font-weight: bold; color: white; background-color: #e53e3e; border-radius: 4px; }}
            .section {{ margin-top: 20px; }}
            .section-title {{ font-size: 16px; font-weight: bold; color: #2b6cb0; border-bottom: 1px solid #cbd5e0; padding-bottom: 4px; }}
            ul {{ line-height: 1.6; }}
            .code-box {{ background: #f7fafc; border: 1px solid #e2e8f0; padding: 10px; font-family: monospace; border-radius: 4px; }}
        </style>
    </head>
    <body>
        <div class="header">
            <div class="title">{title}</div>
            {severity_html}
        </div>
        {summary_section}
        {indicators_section}
        {mitigations_section}
    </body>
    </html>
    """
    
    pdf_path = os.path.join(OUTPUT_DIR, "advisory_report.pdf")
    html_path = os.path.join(OUTPUT_DIR, "advisory_report.html")
    
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html_content)

    try:
        from weasyprint import HTML
        HTML(string=html_content).write_pdf(pdf_path)
        return pdf_path
    except Exception as e:
        print(f"[!] WeasyPrint warning ({e}). HTML report saved to {html_path}")
        return html_path


# ---------------------------------------------------------
# 3. Social Media Generator (LinkedIn & Twitter/X)
# ---------------------------------------------------------
def _social_list(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, (list, tuple, set)):
        return [str(item).strip() for item in value if str(item).strip()]
    text = str(value).strip()
    return [text] if text else []


def select_social_strategy(facts: Dict[str, Any], settings: Dict[str, Any] | None = None) -> Dict[str, str]:
    """Select a source-grounded social archetype and tone using local rules."""
    settings = settings or {}
    severity = str(facts.get("severity", facts.get("risk_level", ""))).upper()
    corpus = " ".join(str(facts.get(key, "")) for key in (
        "title", "summary", "vulnerability", "technical_details", "key_findings", "impact", "recommendations"
    )).lower()
    exploited = any(marker in corpus for marker in (
        "actively exploited", "known exploited", "exploitation in the wild", "cisa kev", "zero-day", "zero day"
    ))
    audience = str(settings.get("audience", settings.get("target_audience", ""))).lower()
    objective = str(settings.get("objective", "")).lower()
    technical_fields = any(_social_list(facts.get(key)) for key in (
        "vulnerability", "technical_details", "attack_vector", "root_cause", "affected_versions"
    ))
    if severity in {"CRITICAL", "HIGH"} or exploited:
        archetype, target_tone = "CRITICAL_ALERT", "Authoritative/Urgent"
    elif technical_fields and (facts.get("cve_ids") or facts.get("cve_id") or facts.get("affected_systems") or facts.get("affected_products")):
        archetype, target_tone = "TECHNICAL_BREAKDOWN", "Analytical"
    elif any(term in audience + " " + objective for term in ("ciso", "leadership", "executive", "governance", "business")):
        archetype, target_tone = "EXECUTIVE_SUMMARY", "Professional"
    else:
        archetype, target_tone = "EDUCATIONAL_INSIGHT", "Instructive"
    return {"content_archetype": archetype, "target_tone": target_tone}


def _social_fact_text(facts: Dict[str, Any], *keys: str, default: str = "") -> str:
    for key in keys:
        values = _social_list(facts.get(key))
        if values:
            return "; ".join(values)
    return default


def _fit_x_post(text: str, limit: int = 279) -> str:
    text = " ".join(text.split())
    if len(text) <= limit:
        return text
    shortened = text[:limit - 1].rsplit(" ", 1)[0].rstrip(" ,;:-")
    return (shortened or text[:limit - 1]).rstrip() + "…"


def _social_hashtags(archetype: str, facts: Dict[str, Any]) -> list[str]:
    """Return a small, relevant, deterministic hashtag set."""
    source_text = " ".join(str(facts.get(key, "")) for key in ("title", "vendor", "source", "affected_systems", "affected_products")).lower()
    tags = ["#CyberSecurity"]
    if archetype == "CRITICAL_ALERT":
        tags.append("#PatchAlert")
        if "cisco" in source_text or any(model in source_text for model in ("rv160", "rv260", "rv340", "rv345")):
            tags.append("#Cisco")
        else:
            tags.append("#PatchManagement")
    elif archetype == "TECHNICAL_BREAKDOWN":
        if "cisco" in source_text or any(model in source_text for model in ("rv160", "rv260", "rv340", "rv345")):
            tags.append("#Cisco")
        tags.append("#VulnerabilityManagement")
    elif archetype == "EXECUTIVE_SUMMARY":
        tags.append("#RiskManagement")
    else:
        tags.append("#SecurityAwareness")
    return tags[:3]


def _normalize_linkedin_hashtags(post: str, archetype: str, facts: Dict[str, Any]) -> str:
    without_tags = re.sub(r"(?<!\w)#[A-Za-z0-9_]+", "", post)
    body = re.sub(r"[ \t]+", " ", without_tags)
    body = re.sub(r"\n(?:[ \t]*\n){2,}", "\n\n", body).strip()
    return f"{body}\n\n{' '.join(_social_hashtags(archetype, facts))}".strip()


def _normalize_x_text(text: str, archetype: str, facts: Dict[str, Any], include_tag: bool = True) -> str:
    without_tags = re.sub(r"(?<!\w)#[A-Za-z0-9_]+", "", text)
    tag = _social_hashtags(archetype, facts)[0] if include_tag else ""
    return _fit_x_post(f"{without_tags.strip()} {tag}".strip())


def _clip_words(value: str, limit: int) -> str:
    words = value.split()
    if len(words) <= limit:
        return value
    return " ".join(words[:limit]).rstrip(" ,;:-") + "…"


def _word_count(value: str) -> int:
    return len(re.findall(r"\b[\w#-]+\b", value, flags=re.UNICODE))


def _structured_linkedin_post(facts: Dict[str, Any], archetype: str) -> str:
    """Build a detailed, operational LinkedIn layout for urgent/technical notices."""
    title = _clip_words(str(facts.get("title") or "Security Advisory").strip(), 14)
    severity = str(facts.get("severity") or facts.get("risk_level") or "Unspecified").strip()
    cves = _social_list(facts.get("cve_ids") or facts.get("cve_id"))
    products = _social_list(facts.get("affected_systems") or facts.get("affected_products"))
    summary = _clip_words(str(facts.get("summary") or "").strip(), 34)
    technical = _clip_words(_social_fact_text(facts, "root_cause", "vulnerability", "technical_details", "attack_vector", default="").strip(), 32)
    interface = _clip_words(_social_fact_text(facts, "affected_interface", "affected_interfaces", "interface", "interface_name", "network_interface", "attack_surface", "management_interface", default="").strip(), 12)
    impacts = _clip_words(_social_fact_text(facts, "impact", "business_impact", default="").strip(), 22)
    source_actions = [_clip_words(action, 22) for action in _social_list(facts.get("recommended_actions") or facts.get("recommendations") or facts.get("mitigation_steps"))[:2]]

    context_parts = [f"Severity is {severity}."]
    if cves:
        context_parts.append(f"Advisory identifier: {', '.join(cves[:4])}.")
    if summary:
        context_parts.append(summary)
    if technical and technical.casefold() not in summary.casefold():
        context_parts.append(f"Root cause / technical detail: {technical}")
    if interface:
        context_parts.append(f"Affected interface or attack surface: {interface}.")
    scope_paragraph = (
        "Network and SOC teams should validate exposure against their own asset inventory. Confirm which listed models and versions are deployed, "
        "whether the relevant interface is enabled and reachable, and which systems require immediate owner review. This post summarizes advisory facts; "
        "it does not assert that a particular organization is exposed or compromised."
    )
    affected_lines = [f"- {product}" for product in products[:8]]
    if len(products) > 8:
        affected_lines.append(f"- Plus {len(products) - 8} additional source-listed products; consult the advisory for the full list.")
    if not affected_lines:
        affected_lines = ["- Product or model names were not identified in the supplied facts; confirm scope in the source advisory."]

    actions = [f"{index}. {action}" for index, action in enumerate(source_actions, start=1)]
    next_number = len(actions) + 1
    actions.extend([
        f"{next_number}. Compare deployed models and versions with the affected scope; identify internet-facing or otherwise reachable management interfaces.",
        f"{next_number + 1}. Apply the vendor-approved patch or mitigation, following the source advisory and local change-control process.",
        f"{next_number + 2}. Where patching is pending, apply source-approved interim controls and restrict relevant access where operationally safe.",
        f"{next_number + 3}. Verify the remediated version, record owner and completion evidence, and review vendor-provided detection guidance with the SOC.",
    ])
    if impacts:
        context_parts.append(f"Source-reported impact: {impacts}.")

    heading = "CRITICAL SECURITY ALERT" if archetype == "CRITICAL_ALERT" else "TECHNICAL SECURITY BREAKDOWN"
    sections = [
        f"{heading}: {title}",
        " ".join(context_parts),
        scope_paragraph,
        "Affected models/products:\n" + "\n".join(affected_lines),
        "Required action plan for Network / IT / SOC teams:\n" + "\n".join(actions),
        "After remediation, retain the affected-asset list, approved change record, validation results, and any escalation notes together. Share status with system owners and incident leadership, and keep unresolved assets visible until their exposure is addressed or formally accepted.",
    ]
    return "\n\n".join(sections)


def _rule_based_social_content(facts: Dict[str, Any], strategy: Dict[str, str]) -> Dict[str, Any]:
    """Build safe fallback copy using only supplied facts and deterministic phrasing."""
    archetype = strategy["content_archetype"]
    title = str(facts.get("title") or "Security update").strip()
    severity = str(facts.get("severity") or facts.get("risk_level") or "Unspecified").strip()
    summary = str(facts.get("summary") or facts.get("vulnerability") or "Review the source advisory for details.").strip()
    cves = _social_list(facts.get("cve_ids") or facts.get("cve_id"))
    affected = _social_list(facts.get("affected_systems") or facts.get("affected_products"))
    actions = _social_list(facts.get("recommended_actions") or facts.get("recommendations") or facts.get("mitigation_steps"))
    impact = _social_fact_text(facts, "impact", "business_impact", default="")
    technical = _social_fact_text(facts, "root_cause", "vulnerability", "technical_details", "attack_vector", default=summary)
    identifiers = ", ".join(cves)
    affected_text = ", ".join(affected)
    action_text = "; ".join(actions)
    hashtags = " ".join(_social_hashtags(archetype, facts))

    if archetype == "CRITICAL_ALERT":
        linkedin = _structured_linkedin_post(facts, archetype)
        x_post = f"Critical security alert: {title}. {identifiers + '. ' if identifiers else ''}Severity {severity}. {action_text or 'Review affected systems and apply source-approved mitigations.'} #CyberSecurity"
    elif archetype == "TECHNICAL_BREAKDOWN":
        linkedin = _structured_linkedin_post(facts, archetype)
        x_post = f"{title}: {technical} {identifiers + '. ' if identifiers else ''}{action_text or 'Follow the advisory for mitigation steps.'} #CyberSecurity"
    elif archetype == "EXECUTIVE_SUMMARY":
        linkedin_parts = [f"Executive security update: {title}", summary]
        if impact: linkedin_parts.append(f"Business impact: {impact}")
        if affected_text: linkedin_parts.append(f"Risk exposure: {affected_text}")
        if action_text: linkedin_parts.append(f"Governance action: {action_text}")
        linkedin = "\n\n".join(linkedin_parts + ["#CyberSecurity #RiskManagement"])
        x_post = f"Executive security update: {title}. {impact or summary} {action_text or 'Review exposure and track remediation.'} #CyberSecurity"
    else:
        linkedin = "\n\n".join([
            f"Security learning: {title}",
            f"{summary}",
            "Key takeaway: Regularly review advisories, affected assets, and recommended safeguards.",
            f"Action to consider: {action_text}" if action_text else "Action to consider: Review your security hygiene and apply relevant safeguards.",
            "#CyberSecurity #SecurityAwareness",
        ])
        x_post = f"Security takeaway: {summary} {action_text or 'Review current advisories and keep safeguards up to date.'} #SecurityAwareness"

    hook = f"Hook: {title}{': ' + identifiers if identifiers else ''}. Severity: {severity}."
    impact_text = _social_fact_text(facts, "impact", "business_impact", default="")
    if not impact_text:
        impact_text = summary if summary not in title else technical
    impact_tweet = f"Impact: {impact_text or 'Review the advisory to understand the affected scope.'}"
    mitigation_tweet = f"Mitigation: {action_text or 'Follow the vendor-approved remediation in the source advisory.'}"
    references = _social_list(facts.get("references") or facts.get("reference_urls"))
    reference_text = ", ".join(references[:2]) or ", ".join(cves) or title + " advisory"
    thread = [
        _fit_x_post(hook),
        _fit_x_post(impact_tweet),
        _fit_x_post(mitigation_tweet),
        _normalize_x_text(f"Reference: {reference_text}", archetype, facts),
    ]
    return {
        "linkedin_post": linkedin,
        "x_post": _normalize_x_text(x_post, archetype, facts),
        "x_thread": thread,
        # Legacy aliases retained for existing GUI callers.
        "linkedin": linkedin,
        "twitter_thread": thread,
    }


def generate_social_content(
    canonical_facts: Dict[str, Any],
    model_name: str = "qwen2.5:3b",
    *,
    strategy: Dict[str, str] | None = None,
    settings: Dict[str, Any] | None = None,
    platform: str = "both",
) -> Dict[str, Any]:
    """Generate archetype-aware social content locally, with a deterministic fallback."""
    settings = settings or {}
    strategy = strategy or select_social_strategy(canonical_facts, settings)
    archetype = strategy.get("content_archetype", "EDUCATIONAL_INSIGHT")
    tone = strategy.get("target_tone", "Professional")
    linkedin_guidance = {
        "CRITICAL_ALERT": (
            "Write an authoritative, structured LinkedIn alert for Network and IT professionals. Start with a high-impact headline. "
            "Give a brief vulnerability summary and root cause; describe Remote Code Execution as root only if that exact fact is in the source. "
            "Add a bulleted affected-model list using only products present in the facts (for example RV160, RV260, RV340, RV345 when listed). "
            "Follow this layout: line 1 urgent headline; one threat-context paragraph with CVE, affected interface, and remote-root-execution risk only when source-supported; "
            "bulleted affected models/products; numbered action plan for IT/SOC teams. Target 150-250 words."
        ),
        "TECHNICAL_BREAKDOWN": (
            "Use the same professional structure: line 1 urgent technical headline; a threat-context paragraph with CVE, affected interface, root cause, "
            "and remote-root-execution risk only when source-supported; bulleted affected models/products; numbered action plan for IT/SOC teams. "
            "Target 150-250 words and explain technical details without adding unsupported claims."
        ),
        "EXECUTIVE_SUMMARY": "Focus on business impact, exposure, and governance recommendations. Avoid unsupported technical claims.",
        "EDUCATIONAL_INSIGHT": "Use an instructive narrative, a clear takeaway, and practical security hygiene guidance grounded in the facts.",
    }.get(archetype, "Write a clear, source-grounded professional update.")
    prompt = f"""
You are a careful cybersecurity social-media editor. Generate content using ONLY the facts below.
Do not invent CVEs, affected products, exploitation status, impact, metrics, links, or mitigations.
Selected archetype: {archetype}
Target tone: {tone}
LinkedIn strategy: {linkedin_guidance}
Audience: {settings.get('audience', 'Security professionals')}
Operator tone preference: {settings.get('operator_tone', settings.get('tone', tone))}

Return JSON with exactly these required keys:
{{"linkedin_post":"...", "x_post":"...", "x_thread":["...", "...", "..."]}}
LinkedIn should follow the selected strategy. For CRITICAL_ALERT and TECHNICAL_BREAKDOWN, produce 150-250 words using the required four-part structured layout.
Place no more than 3 relevant hashtags naturally at the very end.
Do not use generic or irrelevant hashtags, including #malicious. Do not put hashtags in the middle of the post.
x_post must be a single punchy post strictly under 280 characters with relevant hashtags.
x_thread must contain 3 or 4 distinct, non-redundant posts, each strictly under 280 characters, preferably 4, in this order: Hook, Impact, Mitigation, Reference.
Start each tweet with its matching label. For a 3-tweet thread, combine the last two labels as "Mitigation / Reference:". Include a CVE in the hook/reference when available.
Omit unavailable facts rather than guessing; do not create generic filler tweets or repeat the same point.
Return no markdown or surrounding explanation.

FACTS:
{json.dumps(canonical_facts, ensure_ascii=False)}
"""
    try:
        response = ollama.chat(
            model=model_name,
            messages=[{"role": "user", "content": prompt}],
            format="json",
        )
        raw_content = response["message"]["content"]
        social_data = json.loads(raw_content)
        linkedin = str(social_data.get("linkedin_post") or social_data.get("linkedin") or "").strip()
        x_post = str(social_data.get("x_post") or "").strip()
        raw_thread = social_data.get("x_thread") or social_data.get("twitter_thread") or []
        if isinstance(raw_thread, str):
            raw_thread = [raw_thread]
        thread = [str(post).strip() for post in raw_thread if str(post).strip()]
        if not linkedin or not x_post or not 3 <= len(thread) <= 4:
            raise ValueError("Ollama social response did not meet the required post schema or length limits")
        if archetype in {"CRITICAL_ALERT", "TECHNICAL_BREAKDOWN"}:
            if not 150 <= _word_count(linkedin) <= 250:
                raise ValueError("Structured LinkedIn output is outside the 150-250 word target")
            if not re.search(r"(?im)^.*(?:critical security alert|technical security alert|security breakdown).+", linkedin.splitlines()[0]):
                raise ValueError("Structured LinkedIn output is missing its urgent headline")
            if not re.search(r"(?im)^affected (?:models|models/products|products):", linkedin):
                raise ValueError("Structured LinkedIn output is missing the affected-model section")
            if not re.search(r"(?im)^required action plan for (?:network / it / soc|it/soc|network / it) teams:", linkedin):
                raise ValueError("Structured LinkedIn output is missing its numbered action plan")
            affected_models = _social_list(canonical_facts.get("affected_systems") or canonical_facts.get("affected_products"))
            has_numbered_plan = bool(re.search(r"(?m)^\s*1[.)]\s+", linkedin))
            models_are_bulleted = all(
                re.search(r"(?im)^\s*[-*•]\s+.*" + re.escape(model), linkedin)
                for model in affected_models
            )
            if not has_numbered_plan or (affected_models and not models_are_bulleted):
                raise ValueError("Critical LinkedIn output omitted its numbered action plan or affected-model bullets")
        result = {
            "linkedin_post": linkedin,
            "x_post": x_post,
            "x_thread": thread,
            "linkedin": linkedin,
            "twitter_thread": thread,
        }
    except Exception:
        logging.getLogger(__name__).exception("Local social generation failed; using deterministic %s fallback", archetype)
        result = _rule_based_social_content(canonical_facts, strategy)

    result["linkedin_post"] = _normalize_linkedin_hashtags(result.get("linkedin_post", result.get("linkedin", "")), archetype, canonical_facts)
    result["linkedin"] = result["linkedin_post"]
    result["x_post"] = _normalize_x_text(result.get("x_post", ""), archetype, canonical_facts)
    normalized_thread = []
    seen_tweets = set()
    for post in result.get("x_thread", result.get("twitter_thread", [])):
        clean_post = _normalize_x_text(str(post), archetype, canonical_facts, include_tag=False)
        fingerprint = " ".join(clean_post.split()).casefold()
        if clean_post and fingerprint not in seen_tweets:
            normalized_thread.append(clean_post)
            seen_tweets.add(fingerprint)
    if len(normalized_thread) < 3:
        result = _rule_based_social_content(canonical_facts, strategy)
        result["linkedin_post"] = _normalize_linkedin_hashtags(result["linkedin_post"], archetype, canonical_facts)
        result["linkedin"] = result["linkedin_post"]
        normalized_thread = [
            _normalize_x_text(str(post), archetype, canonical_facts, include_tag=False)
            for post in result["x_thread"]
        ]
    if normalized_thread:
        roles = ["Hook", "Impact", "Mitigation / Reference"] if len(normalized_thread) == 3 else ["Hook", "Impact", "Mitigation", "Reference"]
        labeled_thread = []
        for post, role in zip(normalized_thread, roles):
            content = re.sub(r"^(?:Hook|Impact|Mitigation(?: / Reference)?|Reference)\s*:\s*", "", post, flags=re.IGNORECASE)
            labeled_thread.append(_normalize_x_text(f"{role}: {content}", archetype, canonical_facts, include_tag=False))
        normalized_thread = labeled_thread
        normalized_thread[-1] = _normalize_x_text(normalized_thread[-1], archetype, canonical_facts)
    result["x_thread"] = normalized_thread[:4]
    result["twitter_thread"] = result["x_thread"]

    if archetype in {"CRITICAL_ALERT", "TECHNICAL_BREAKDOWN"} and not 150 <= _word_count(result["linkedin_post"]) <= 250:
        result = _rule_based_social_content(canonical_facts, strategy)
        result["linkedin_post"] = _normalize_linkedin_hashtags(result["linkedin_post"], archetype, canonical_facts)
        result["linkedin"] = result["linkedin_post"]
        result["x_post"] = _normalize_x_text(result.get("x_post", ""), archetype, canonical_facts)
        result["x_thread"] = result["x_thread"][:4]
        result["twitter_thread"] = result["x_thread"]

    selected = str(platform or "both").lower()
    if selected == "linkedin":
        result.update({"x_post": "", "x_thread": [], "twitter_thread": []})
    elif selected in {"x", "twitter"}:
        result.update({"linkedin_post": "", "linkedin": ""})
    try:
        output_path = os.path.join(OUTPUT_DIR, "social_posts.json")
        with open(output_path, "w", encoding="utf-8") as f:
            json.dump(result, f, ensure_ascii=False, indent=2)
    except OSError:
        logging.getLogger(__name__).warning("Could not save generated social_posts.json", exc_info=True)
    return result


# ---------------------------------------------------------
# 4. PowerPoint Slide Deck Generator (.pptx)
# ---------------------------------------------------------
def generate_presentation(canonical_facts: Dict[str, Any]) -> str:
    prs = Presentation()
    
    # Slide 1: Title Slide
    blank_slide_layout = prs.slide_layouts[0]
    slide1 = prs.slides.add_slide(blank_slide_layout)
    slide1.shapes.title.text = canonical_facts.get("title", "Threat Briefing")
    slide1.placeholders[1].text = f"Severity: {canonical_facts.get('severity', 'HIGH')}\nNTRO Cybersecurity Automated Briefing"
    
    # Slide 2: Threat Summary
    bullet_slide_layout = prs.slide_layouts[1]
    slide2 = prs.slides.add_slide(bullet_slide_layout)
    slide2.shapes.title.text = "Executive Threat Overview"
    tf2 = slide2.placeholders[1].text_frame
    tf2.text = canonical_facts.get("summary", "")
    p2 = tf2.add_paragraph()
    p2.text = f"Affected Systems: {', '.join(canonical_facts.get('affected_systems', []))}"
    
    # Slide 3: Vulnerabilities & IPs
    slide3 = prs.slides.add_slide(bullet_slide_layout)
    slide3.shapes.title.text = "Technical Indicators & CVEs"
    tf3 = slide3.placeholders[1].text_frame
    tf3.text = f"Identified CVEs: {', '.join(canonical_facts.get('cve_ids', []))}"
    p3 = tf3.add_paragraph()
    p3.text = f"Flagged IP Addresses: {', '.join(canonical_facts.get('locked_ips', []))}"
    
    # Slide 4: Recommendations
    slide4 = prs.slides.add_slide(bullet_slide_layout)
    slide4.shapes.title.text = "Mitigation & Next Steps"
    tf4 = slide4.placeholders[1].text_frame
    for action in canonical_facts.get("recommended_actions", []):
        p = tf4.add_paragraph()
        p.text = f"• {action}"

    pptx_path = os.path.join(OUTPUT_DIR, "presentation.pptx")
    prs.save(pptx_path)
    return pptx_path


# ---------------------------------------------------------
# 5. Video Package Generator (Script + Storyboard)
# ---------------------------------------------------------
def generate_video_package(canonical_facts: Dict[str, Any], model_name: str = "qwen2.5:3b") -> Dict[str, Any]:
    prompt = f"""
    Create a 30-second cybersecurity alert video script and visual storyboard based on these facts:
    {json.dumps(canonical_facts)}
    
    Return a JSON object with keys:
    - "video_title": str
    - "scenes": list of objects containing ("scene_number", "visual_description", "narration_text", "on_screen_text")
    """
    
    response = ollama.chat(
        model=model_name,
        messages=[{'role': 'user', 'content': prompt}],
        format='json'
    )
    
    video_data = json.loads(response['message']['content'])
    output_path = os.path.join(OUTPUT_DIR, "video_package.json")
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(video_data, f, indent=2)
        
    return video_data


# ---------------------------------------------------------
# 6. Infographic Blueprint Generator (.svg)
# ---------------------------------------------------------
def _svg_clip(text: str, limit: int) -> str:
    """Collapse whitespace and clip SVG label text with an explicit ellipsis."""
    text = " ".join(str(text).split())
    return text if len(text) <= limit else text[: limit - 1].rstrip(" ,;:-") + "…"


def generate_infographic_svg(canonical_facts: Dict[str, Any]) -> str:
    """Render a responsive SVG infographic strictly derived from the supplied facts."""
    facts = canonical_facts or {}
    design = infer_design(facts)
    c = design["palette"]
    layout = design["layout"]
    title = escape(_fact_headline(facts))
    cves = _fact_items(facts.get("cve_ids") or facts.get("cve_id"))
    ips = _fact_items(facts.get("locked_ips"))
    systems = _fact_items(facts.get("affected_systems") or facts.get("affected_products"))
    actions = _fact_items(
        facts.get("recommended_actions") or facts.get("recommendations") or facts.get("mitigation_steps") or facts.get("key_points")
    )
    severities = _fact_items(facts.get("severity") or facts.get("risk_level"))
    severity = severities[0].upper() if severities else ""
    severity_color = _SEVERITY_COLORS.get(severity, c["secondary"])
    summary_text = " ".join(_fact_items(facts.get("summary")))
    summary_lines = wrap(summary_text, 100) if summary_text else []
    if len(summary_lines) > 6:
        summary_lines = summary_lines[:5] + [_svg_clip(" ".join(summary_lines[5:]), 100)]

    badge_html = ""
    if severity:
        badge_html = (
            f'<rect x="800" y="36" width="155" height="39" rx="19" fill="{severity_color}"/>'
            f'<text x="877" y="61" text-anchor="middle" fill="#fff" font-family="Arial" font-size="14" font-weight="700">{escape(_svg_clip(severity, 16))}</text>'
        )

    body = []
    y = 178

    # Identifier chips: one chip per CVE / IP actually present in the source facts.
    badges = [*(f"CVE {value}" for value in cves[:4]), *(f"IP {value}" for value in ips[:3])]
    if badges:
        x = 40
        for badge in badges:
            label = escape(str(badge)[:34])
            chip_width = min(230, max(110, len(label) * 8 + 32))
            if x + chip_width > 965:
                break
            body.append(f'<rect x="{x}" y="{y}" width="{chip_width}" height="30" rx="15" fill="{c["surface"]}" stroke="{c["highlight"]}"/><text x="{x+14}" y="{y+20}" fill="{c["ink"]}" font-family="Arial" font-size="12">{label}</text>')
            x += chip_width + 10
        y += 46

    # Brief: only the supplied summary, no filler copy.
    if summary_lines:
        body.append(f'<text x="40" y="{y+14}" fill="{c["muted"]}" font-family="Arial" font-size="13" font-weight="700">BRIEF</text>')
        for index, line in enumerate(summary_lines):
            body.append(f'<text x="40" y="{y+40+index*22}" fill="{c["ink"]}" font-family="Arial" font-size="15">{escape(line)}</text>')
        y += 40 + 22 * len(summary_lines) + 14

    # Key callouts: only fact categories that exist in the input are rendered.
    callouts = []
    if severity:
        callouts.append(("SEVERITY RATING", severity, "alert", severity_color))
    if systems:
        callouts.append(("AFFECTED SYSTEMS", ", ".join(systems), "server", c["highlight"]))
    if cves:
        callouts.append(("IDENTIFIED CVEs", ", ".join(cves), "shield", c["accent"]))
    if ips:
        callouts.append(("TRACKED IPs", ", ".join(ips), "lock", c["secondary"]))
    if callouts:
        columns = min(3, len(callouts))
        card_width = (920 - (columns - 1) * 14) // columns
        wrap_chars = max(14, min(60, int((card_width - 48) / 7.3)))
        for index, (label, value, icon_name, accent) in enumerate(callouts):
            row, column = divmod(index, columns)
            card_x, card_top = 40 + column * (card_width + 14), y + row * 114
            body.append(f'<rect x="{card_x}" y="{card_top}" width="{card_width}" height="100" rx="14" fill="{c["surface"]}" stroke="{c["border"]}" stroke-width="1" filter="url(#shadow)"/><rect x="{card_x}" y="{card_top}" width="5" height="100" rx="2" fill="{accent}"/><rect x="{card_x+1}" y="{card_top+1}" width="{card_width-2}" height="98" rx="13" fill="url(#accentWash)" opacity=".45"/><text x="{card_x+24}" y="{card_top+26}" fill="{c["muted"]}" font-family="Arial" font-size="11" font-weight="700">{escape(label)}</text>')
            body.append(svg_icon(icon_name, card_x + card_width - 40, card_top + 12, accent, 22))
            value_lines = wrap(value, wrap_chars) or [value]
            if len(value_lines) > 3:
                value_lines = value_lines[:2] + [_svg_clip(" ".join(value_lines[2:]), wrap_chars)]
            for line_index, line in enumerate(value_lines):
                body.append(f'<text x="{card_x+24}" y="{card_top+52+line_index*17}" fill="{c["ink"]}" font-family="Arial" font-size="14">{escape(line)}</text>')
        callout_rows = (len(callouts) + columns - 1) // columns
        y += callout_rows * 114 - 14 + 18

    # Signal profile: bar lengths are counts of the fact entries above, nothing more.
    signals = [
        ("CVEs", len(cves), c["accent"]),
        ("IP indicators", len(ips), c["highlight"]),
        ("Affected systems", len(systems), c["secondary"]),
        ("Response steps", len(actions), c["success"]),
    ]
    signals = [entry for entry in signals if entry[1] > 0]
    if signals:
        max_count = max(entry[1] for entry in signals)
        body.append(f'<text x="62" y="{y+24}" fill="{c["muted"]}" font-family="Arial" font-size="13" font-weight="700">SOURCE SIGNAL PROFILE</text>')
        for index, (label, count, color) in enumerate(signals):
            row_y = y + 44 + index * 38
            bar_width = max(26, int(560 * count / max_count))
            body.append(f'<text x="62" y="{row_y+9}" fill="{c["ink"]}" font-family="Arial" font-size="12">{escape(label)}</text>')
            body.append(f'<rect x="250" y="{row_y}" width="560" height="10" rx="5" fill="{c["border"]}"/><rect x="250" y="{row_y}" width="{bar_width}" height="10" rx="5" fill="{color}"/><text x="{262+bar_width}" y="{row_y+9}" fill="{c["ink"]}" font-family="Arial" font-size="11" font-weight="700">{count}</text>')
        y += 44 + len(signals) * 38 + 14

    # Severity scale: only when the rating matches the fixed taxonomy.
    if severity in _SEVERITY_COLORS:
        ramp = [("LOW", "#34866D"), ("MEDIUM", "#D28A25"), ("HIGH", "#D14B3F"), ("CRITICAL", "#B42332")]
        segment_width = 176
        body.append(f'<text x="62" y="{y+26}" fill="{c["muted"]}" font-family="Arial" font-size="13" font-weight="700">SEVERITY LEVEL</text>')
        for index, (level, color) in enumerate(ramp):
            segment_x = 178 + index * segment_width
            active = level == severity
            body.append(f'<rect x="{segment_x+4}" y="{y+14}" width="{segment_width-8}" height="12" rx="6" fill="{color}" opacity="{1 if active else 0.25}"/>')
            body.append(f'<text x="{segment_x+segment_width//2}" y="{y+46}" text-anchor="middle" fill="{c["ink"] if active else c["muted"]}" font-family="Arial" font-size="10" font-weight="{700 if active else 400}">{level}</text>')
            if active:
                body.append(f'<circle cx="{segment_x+segment_width//2}" cy="{y+8}" r="5" fill="{color}" stroke="{c["surface"]}" stroke-width="2"/>')
        y += 60

    # Response panel: rendered only when recommended actions exist.
    if actions:
        action_lines = [(str(action), wrap(str(action), 64) or [""]) for action in actions]
        if layout == "timeline":
            panel_height = max(180, 78 + len(actions) * 66)
        elif layout == "process_flow":
            panel_height = max(190, 72 + ((len(actions) + 2) // 3) * 112)
        elif layout == "split_comparison":
            panel_height = max(190, 125 + ((len(actions) + 1) // 2) * 38)
        else:
            panel_height = max(180, 65 + ((len(actions) + 1) // 2) * 44)
        body.append(f'<rect x="40" y="{y}" width="920" height="{panel_height}" rx="16" fill="{c["surface"]}" stroke="{c["border"]}" stroke-width="1" filter="url(#shadow)"/><text x="86" y="{y+35}" fill="{c["primary"]}" font-family="Arial" font-size="17" font-weight="700">RECOMMENDED RESPONSE</text><rect x="62" y="{y+53}" width="876" height="2" fill="{c["border"]}"/><rect x="62" y="{y+53}" width="136" height="3" rx="2" fill="{c["accent"]}"/>')
        body.append(svg_icon("check", 62, y+17, c["success"], 22))
        if layout == "process_flow":
            for index, (action_text, _) in enumerate(action_lines):
                row, column = divmod(index, 3)
                box_x, box_top = 62 + column * 290, y + 66 + row * 112
                body.append(f'<rect x="{box_x}" y="{box_top}" width="260" height="88" rx="12" fill="{c["background"]}" stroke="{c["highlight"]}"/><circle cx="{box_x+25}" cy="{box_top+25}" r="15" fill="{c["primary"]}"/><text x="{box_x+25}" y="{box_top+30}" text-anchor="middle" fill="#fff" font-family="Arial" font-size="13">{index+1}</text>')
                flow_lines = wrap(action_text, 34)
                if len(flow_lines) > 3:
                    flow_lines = flow_lines[:2] + [_svg_clip(" ".join(flow_lines[2:]), 34)]
                for line_index, line in enumerate(flow_lines):
                    body.append(f'<text x="{box_x+50}" y="{box_top+29+line_index*17}" fill="{c["ink"]}" font-family="Arial" font-size="12">{escape(line)}</text>')
        elif layout == "timeline":
            body.append(f'<path d="M84 {y+73}v{max(40, len(actions)*66)}" stroke="{c["highlight"]}" stroke-width="4"/>')
            for index, (_, lines) in enumerate(action_lines):
                line_y = y + 78 + index * 66
                body.append(f'<circle cx="84" cy="{line_y-5}" r="10" fill="{c["accent"]}"/><text x="112" y="{line_y}" fill="{c["ink"]}" font-family="Arial" font-size="13">{index+1:02d}  {escape(_svg_clip(" ".join(lines), 105))}</text>')
        elif layout == "split_comparison":
            split = max(1, (len(actions) + 1) // 2)
            for column, heading, items in ((0, "ASSESSMENT", action_lines[:split]), (1, "RESPONSE", action_lines[split:] or action_lines[:1])):
                box_x = 62 + column * 440
                body.append(f'<rect x="{box_x}" y="{y+65}" width="418" height="{panel_height-82}" rx="12" fill="{c["background"]}"/><text x="{box_x+18}" y="{y+92}" fill="{c["secondary"]}" font-family="Arial" font-size="12" font-weight="700">{heading}</text>')
                for index, (_, lines) in enumerate(items):
                    body.append(f'<text x="{box_x+20}" y="{y+122+index*35}" fill="{c["ink"]}" font-family="Arial" font-size="12">• {escape(_svg_clip(" ".join(lines), 60))}</text>')
        else:
            columns = 2 if len(actions) > 2 else 1
            box_width = (876 - (columns - 1) * 14) // columns
            row_limit = max(40, int((box_width - 40) / 6.2))
            for index, (_, lines) in enumerate(action_lines):
                row, column = divmod(index, columns)
                box_x, box_top = 62 + column * (box_width + 14), y + 64 + row * 44
                body.append(f'<rect x="{box_x}" y="{box_top}" width="{box_width}" height="34" rx="9" fill="{c["background"]}"/><circle cx="{box_x+17}" cy="{box_top+17}" r="6" fill="{c["accent"]}"/><text x="{box_x+32}" y="{box_top+22}" fill="{c["ink"]}" font-family="Arial" font-size="12">{escape(_svg_clip(" ".join(lines), row_limit))}</text>')
        y += panel_height + 16

    height = max(y + 18, 260)
    parts = [f'''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1000 {height}" width="100%" role="img" aria-label="{title}">
<defs><filter id="shadow" x="-20%" y="-20%" width="140%" height="150%"><feDropShadow dx="0" dy="5" stdDeviation="7" flood-color="{c['primary']}" flood-opacity=".14"/></filter><linearGradient id="banner"><stop stop-color="{c['primary']}"/><stop offset="1" stop-color="{c['secondary']}"/></linearGradient><linearGradient id="accentWash" x2="1" y2="1"><stop stop-color="{c['accent']}" stop-opacity=".18"/><stop offset="1" stop-color="{c['highlight']}" stop-opacity=".04"/></linearGradient></defs><style>text{{font-family:'Segoe UI',Inter,Arial,sans-serif;line-height:1.45}}</style>
<rect width="100%" height="100%" rx="22" fill="{c['background']}"/><rect width="1000" height="158" rx="22" fill="url(#banner)"/><path d="M0 138 Q500 176 1000 132 V158 Q1000 176 980 176 H20 Q0 176 0 158Z" fill="{c['accent']}"/>
{svg_icon(design['icon'], 40, 35, '#FFFFFF', 32)}<text x="88" y="58" fill="#fff" font-family="Arial" font-size="13" font-weight="700" letter-spacing="2">{escape(design['topic'].upper())} / {layout.replace('_',' ').upper()}</text>
<text x="40" y="108" fill="#fff" font-family="Arial" font-size="{max(16,min(32,int(920/max(1,len(title)*.56))))}" font-weight="700">{title}</text>{badge_html}''']
    parts.extend(body)
    parts.append('</svg>')
    svg_path = os.path.join(OUTPUT_DIR, "infographic_blueprint.svg")
    with open(svg_path, "w", encoding="utf-8") as f:
        f.write("\n".join(parts))
    return svg_path
