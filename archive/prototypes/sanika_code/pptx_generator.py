from pptx import Presentation
import json
import os


# ============================================
# STEP 1: Load presentation JSON
# ============================================

json_path = "output/presentation.json"

with open(json_path, "r", encoding="utf-8") as file:
    presentation_data = json.load(file)


# ============================================
# STEP 2: Create PowerPoint
# ============================================

prs = Presentation()


# ============================================
# STEP 3: Add title slide
# ============================================

title_slide = prs.slides.add_slide(
    prs.slide_layouts[0]
)

title_slide.shapes.title.text = presentation_data[
    "presentation_title"
]

title_slide.placeholders[1].text = (
    "Generated using AI"
)


# ============================================
# STEP 4: Add content slides
# ============================================

for slide_data in presentation_data["slides"]:

    slide = prs.slides.add_slide(
        prs.slide_layouts[1]
    )

    # Slide title
    slide.shapes.title.text = slide_data["title"]

    # Content placeholder
    content = slide.placeholders[1]

    content.text = ""

    # Add bullet points
    for point in slide_data["key_points"]:

        paragraph = content.text_frame.add_paragraph()

        paragraph.text = point

        paragraph.level = 0


# ============================================
# STEP 5: Save PowerPoint
# ============================================

output_path = "output/presentation.pptx"

prs.save(output_path)


print("\n============================================")
print("SUCCESS!")
print("============================================")

print("\nPowerPoint created successfully!")

print(f"File location: {output_path}")